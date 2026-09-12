"""扫码参保：提交审核队列（管理端，带鉴权）。

跟 routers/enroll.py（免登录公开提交）刻意分成两个文件：那边一个登录态都
不认，这边全部要鉴权——混在一个文件里容易在改动时把公开/私有的边界改糊。

审核是强制门：PositionEnrollSubmission 只有走 approve 才落地为正式
InsuredPerson（解密身份证号写入明文字段，跟 insured.py 的 add_person 建的
是同一种记录）。按本轮范围收窄的决定不做使用费余额校验（require_usage_funded
留给下一步）。

- 平台（admin）：看全部、审核、导出
- 企业端（enterprise）：只能看自己企业岗位下的提交，不能审核
"""
import io
from datetime import datetime, timezone
from typing import Optional

import openpyxl
from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..core.audit import audit
from ..core.db import db
from ..core.id_number import id_decrypt, mask_id_number
from ..core.rbac import require_role
from ..core.security import current_user
from ..models import (
    Enterprise, InsurancePlan, InsuredPerson, PositionEnrollSubmission,
    User, WorkPosition,
)
from ..schemas import EnrollReviewIn

router = APIRouter(prefix="/api", tags=["enroll-review"])

_REVIEW_LABEL = {"pending": "待审核", "approved": "已通过", "rejected": "已拒绝"}
_PAYMENT_LABEL = {"not_required": "无需支付", "pending": "待支付", "paid": "已支付"}
_MODE_LABEL = {"personal": "个人缴纳", "employer": "单位缴纳"}


def _submission_payload(session: Session, item: PositionEnrollSubmission) -> dict:
    position = session.get(WorkPosition, item.position_id)
    enterprise = session.get(Enterprise, position.enterprise_id) if position else None
    plan = session.get(InsurancePlan, position.plan_id) if position and position.plan_id else None
    return {
        "id": item.id,
        "position_id": item.position_id,
        "position_name": position.name if position else "",
        "actual_employer": position.actual_employer if position else "",
        "enterprise_id": enterprise.id if enterprise else None,
        "enterprise_name": enterprise.name if enterprise else "",
        "plan_name": plan.name if plan else "",
        "payment_mode": item.payment_mode,
        "payment_mode_label": _MODE_LABEL.get(item.payment_mode, item.payment_mode),
        "name": item.name,
        # 列表只给脱敏号码；审核员核对全号时用导出（有 audit 记录）
        "id_number_masked": mask_id_number(id_decrypt(item.id_number_cipher)) if item.id_number_cipher else "",
        "phone": item.phone,
        "payment_status": item.payment_status,
        "payment_status_label": _PAYMENT_LABEL.get(item.payment_status, item.payment_status),
        "order_no": item.order_no,
        "review_status": item.review_status,
        "review_status_label": _REVIEW_LABEL.get(item.review_status, item.review_status),
        "review_note": item.review_note,
        "reviewed_at": item.reviewed_at,
        "insured_person_id": item.insured_person_id,
        "created_at": item.created_at,
    }


def _scoped_statement(user: User, review_status: Optional[str], position_id: Optional[int]):
    stmt = select(PositionEnrollSubmission).order_by(PositionEnrollSubmission.id.desc())
    if review_status:
        stmt = stmt.where(PositionEnrollSubmission.review_status == review_status)
    if position_id:
        stmt = stmt.where(PositionEnrollSubmission.position_id == position_id)
    if user.role == "enterprise":
        if not user.enterprise_id:
            return stmt.where(PositionEnrollSubmission.id.is_(None))
        stmt = stmt.join(WorkPosition, PositionEnrollSubmission.position_id == WorkPosition.id)
        stmt = stmt.where(WorkPosition.enterprise_id == user.enterprise_id)
    elif user.role != "admin":
        raise HTTPException(403, "无权查看参保提交")
    return stmt


@router.get("/enroll-submissions")
def enroll_submissions(
    review_status: Optional[str] = Query(None),
    position_id: Optional[int] = Query(None),
    user: User = Depends(current_user),
    session: Session = Depends(db),
):
    stmt = _scoped_statement(user, review_status, position_id)
    return [_submission_payload(session, x) for x in session.scalars(stmt)]


@router.patch("/enroll-submissions/{item_id}/review", dependencies=[Depends(require_role("admin", detail="仅平台端可审核参保提交"))])
def review_enroll_submission(item_id: int, data: EnrollReviewIn, user: User = Depends(current_user), session: Session = Depends(db)):
    item = session.get(PositionEnrollSubmission, item_id)
    if not item:
        raise HTTPException(404, "提交记录不存在")
    if item.review_status != "pending":
        raise HTTPException(400, f"该提交已是「{_REVIEW_LABEL.get(item.review_status, item.review_status)}」，不能重复审核")
    if data.status == "rejected":
        item.review_status = "rejected"
        item.review_note = data.review_note
        item.reviewed_by = user.id
        item.reviewed_at = datetime.now(timezone.utc)
        session.commit()
        audit(session, user, "reject", "enroll_submission", str(item.id), data.review_note)
        return _submission_payload(session, item)

    # 通过：个人缴纳必须已支付才能落地（钱没到就把人转成正式参保，
    # 后面追缴没有任何抓手）；单位缴纳本轮不查余额，直接放行。
    if item.payment_mode == "personal" and item.payment_status != "paid":
        raise HTTPException(400, "个人缴纳提交需支付完成后才能审核通过")
    position = session.get(WorkPosition, item.position_id)
    if not position or position.status != "approved":
        raise HTTPException(400, "关联岗位不存在或未审核通过")

    person = InsuredPerson(
        enterprise_id=position.enterprise_id,
        position_id=position.id,
        name=item.name,
        id_number=id_decrypt(item.id_number_cipher) if item.id_number_cipher else "",
        phone=item.phone,
        occupation=position.name,
        occupation_class=position.occupation_class,
        status="pending",
    )
    session.add(person)
    session.flush()
    item.review_status = "approved"
    item.review_note = data.review_note
    item.reviewed_by = user.id
    item.reviewed_at = datetime.now(timezone.utc)
    item.insured_person_id = person.id
    session.commit()
    audit(session, user, "approve", "enroll_submission", str(item.id), f"insured_person={person.id}")
    return _submission_payload(session, item)


@router.get("/enroll-submissions/export")
def export_enroll_submissions(
    review_status: Optional[str] = Query(None),
    position_id: Optional[int] = Query(None),
    user: User = Depends(current_user),
    session: Session = Depends(db),
):
    stmt = _scoped_statement(user, review_status, position_id)
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "扫码参保提交"
    sheet.append(["姓名", "身份证号", "手机号", "缴费方式", "支付状态", "审核状态", "审核备注", "投保单位", "岗位", "提交时间"])
    count = 0
    for x in session.scalars(stmt):
        payload = _submission_payload(session, x)
        # 平台导出给全号（审核员要拿去和保司核对名单）；企业端只给脱敏号
        id_display = id_decrypt(x.id_number_cipher) if (user.role == "admin" and x.id_number_cipher) else payload["id_number_masked"]
        sheet.append([
            x.name, id_display, x.phone or "",
            payload["payment_mode_label"], payload["payment_status_label"],
            payload["review_status_label"], x.review_note or "",
            payload["enterprise_name"], payload["position_name"],
            x.created_at.strftime("%Y-%m-%d %H:%M") if x.created_at else "",
        ])
        count += 1
    for row_number in range(2, sheet.max_row + 1):
        sheet.cell(row_number, 2).number_format = "@"
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = min(32, max(10, max(len(str(cell.value or "")) for cell in column) + 2))
    output = io.BytesIO()
    book.save(output)
    output.seek(0)
    audit(session, user, "export", "enroll_submission", f"count={count}")
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": "attachment; filename=enroll-submissions.xlsx"})
